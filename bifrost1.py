import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from PIL import Image, ImageTk, ImageGrab
import pytesseract
from docx import Document
import cv2
import numpy as np
import pyautogui
import keyboard
import pyperclip
import sys
from pathlib import Path
from datetime import datetime

BASE_DIR = Path(r"D:\Bifrost 1")
SCREENSHOT_DIR = BASE_DIR / "Screenshots"
TEXT_DIR = BASE_DIR / "Extracted_Text"
CONFIG_FILE = BASE_DIR / "config.cfg"

for directory in [BASE_DIR, SCREENSHOT_DIR, TEXT_DIR]:
    directory.mkdir(parents=True, exist_ok=True)

class ScreenshotOverlay(tk.Toplevel):
    def __init__(self, parent):
        super().__init__(parent)
        self.attributes('-fullscreen', True, '-alpha', 0.3)
        self.canvas = tk.Canvas(self, cursor="cross", highlightthickness=0)
        self.canvas.pack(fill="both", expand=True)
        
        self.start_x = None
        self.start_y = None
        self.rect = None
        self.selection = None
        
        self.canvas.bind("<Button-1>", self.start_select)
        self.canvas.bind("<B1-Motion>", self.update_select)
        self.canvas.bind("<ButtonRelease-1>", self.finalize_select)
        self.bind("<Escape>", self.cancel_select)

    def start_select(self, event):
        self.start_x = self.canvas.canvasx(event.x)
        self.start_y = self.canvas.canvasy(event.y)
        self.rect = self.canvas.create_rectangle(
            self.start_x, self.start_y, self.start_x, self.start_y,
            outline="#ff0000", width=2, tags="selection"
        )

    def update_select(self, event):
        cur_x = self.canvas.canvasx(event.x)
        cur_y = self.canvas.canvasy(event.y)
        self.canvas.coords(self.rect, self.start_x, self.start_y, cur_x, cur_y)

    def finalize_select(self, event):
        x1, y1, x2, y2 = self.canvas.coords(self.rect)
        self.selection = (
            min(x1, x2), min(y1, y2),
            max(x1, x2), max(y1, y2)
        )
        self.destroy()

    def cancel_select(self, event=None):
        self.selection = None
        self.destroy()

class BifrostApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Bifrost 1 - Text Extraction Suite")
        self.root.geometry("1100x800")
        self.setup_gui()
        self.load_config()

        self.original_image = None
        self.current_image = None
        self.extracted_text = ""
        self.image_scale = 1.0
        self.image_position = [0, 0]
        self.drag_start = None
        
        keyboard.add_hotkey('ctrl+shift+s', self.capture_screenshot)
        self.root.protocol("WM_DELETE_WINDOW", self.on_close)

    def setup_gui(self):

        self.style = ttk.Style()
        self.style.configure("TButton", padding=6, font=('Segoe UI', 10))
        self.style.configure("Title.TLabel", font=('Segoe UI', 18, 'bold'))
        self.style.configure("Status.TLabel", font=('Segoe UI', 9), foreground="#666")

        main_frame = ttk.Frame(self.root)
        main_frame.pack(fill="both", expand=True, padx=20, pady=20)

        header_frame = ttk.Frame(main_frame)
        header_frame.pack(fill="x", pady=(0, 20))
        ttk.Label(header_frame, 
                text="Bifrost 1 - Text Extraction Suite",
                style="Title.TLabel").pack(side=tk.LEFT)

        btn_frame = ttk.Frame(main_frame)
        btn_frame.pack(fill="x", pady=10)
        
        ttk.Button(btn_frame, 
                 text="Capture Screenshot (Ctrl+Shift+S)",
                 command=self.capture_screenshot).pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, 
                 text="Upload Image",
                 command=self.upload_image).pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame,
                 text="Configure Tesseract",
                 command=self.configure_tesseract).pack(side=tk.RIGHT, padx=5)

        self.canvas = tk.Canvas(main_frame, bg="#ffffff", relief="solid")
        self.canvas.pack(pady=10, fill="both", expand=True)

        zoom_frame = ttk.Frame(main_frame)
        zoom_frame.pack(pady=5)
        ttk.Button(zoom_frame, 
                 text="Reset Zoom", 
                 command=self.reset_zoom).pack(side=tk.LEFT, padx=5)
        self.zoom_label = ttk.Label(zoom_frame, text="Zoom: 100%")
        self.zoom_label.pack(side=tk.LEFT, padx=5)

        text_frame = ttk.Frame(main_frame)
        text_frame.pack(fill="both", expand=True)
        self.txt_display = tk.Text(text_frame,
                                 wrap=tk.WORD,
                                 font=('Consolas', 10),
                                 padx=10,
                                 pady=10)
        scrollbar = ttk.Scrollbar(text_frame,
                                orient="vertical",
                                command=self.txt_display.yview)
        self.txt_display.configure(yscrollcommand=scrollbar.set)
        scrollbar.pack(side=tk.RIGHT, fill="y")
        self.txt_display.pack(fill="both", expand=True)

        self.status = ttk.Label(main_frame,
                              text="Ready",
                              style="Status.TLabel")
        self.status.pack(fill="x", pady=(10, 0))

        export_frame = ttk.Frame(main_frame)
        export_frame.pack(pady=10)
        ttk.Button(export_frame,
                 text="Copy to Clipboard",
                 command=self.copy_clipboard).pack(side=tk.LEFT, padx=5)
        ttk.Button(export_frame,
                 text="Save as Text",
                 command=lambda: self.save_text()).pack(side=tk.LEFT, padx=5)
        ttk.Button(export_frame,
                 text="Save as Word",
                 command=lambda: self.save_word()).pack(side=tk.LEFT, padx=5)

        self.canvas.bind("<MouseWheel>", self.on_mousewheel)
        self.canvas.bind("<ButtonPress-1>", self.start_pan)
        self.canvas.bind("<B1-Motion>", self.pan_image)
        self.canvas.bind("<ButtonRelease-1>", self.end_pan)

    def show_preview(self, image_path):
        try:
            self.original_image = Image.open(image_path)
            self.current_image = self.original_image.copy()
            self.image_scale = 1.0
            self.image_position = [0, 0]
            self.update_zoom_display()
            self.update_canvas_image()
        except Exception as e:
            self.show_error(f"Preview Error: {str(e)}")

    def update_canvas_image(self):
        width = int(self.original_image.width * self.image_scale)
        height = int(self.original_image.height * self.image_scale)
        
        resized_img = self.original_image.resize(
            (width, height), 
            resample=Image.Resampling.LANCZOS
        )
        
        self.tk_image = ImageTk.PhotoImage(resized_img)
        self.canvas.delete("all")
        self.canvas.create_image(
            self.image_position[0], 
            self.image_position[1], 
            anchor="nw", 
            image=self.tk_image
        )
        self.canvas.config(scrollregion=self.canvas.bbox("all"))

    def on_mousewheel(self, event):
        zoom_factor = 1.1 if event.delta > 0 else 0.9
        new_scale = self.image_scale * zoom_factor
        
        if 0.1 <= new_scale <= 5.0:
            self.image_scale = new_scale
            x = self.canvas.canvasx(event.x)
            y = self.canvas.canvasy(event.y)
            
            self.image_position[0] = x - (x - self.image_position[0]) * zoom_factor
            self.image_position[1] = y - (y - self.image_position[1]) * zoom_factor
            
            self.update_zoom_display()
            self.update_canvas_image()

    def start_pan(self, event):
        self.drag_start = (event.x, event.y)
        self.canvas.config(cursor="fleur")

    def pan_image(self, event):
        if self.drag_start:
            dx = event.x - self.drag_start[0]
            dy = event.y - self.drag_start[1]
            
            self.image_position[0] += dx
            self.image_position[1] += dy
            
            self.drag_start = (event.x, event.y)
            self.update_canvas_image()

    def end_pan(self, event):
        self.drag_start = None
        self.canvas.config(cursor="")

    def reset_zoom(self):
        self.image_scale = 1.0
        self.image_position = [0, 0]
        self.update_zoom_display()
        self.update_canvas_image()

    def update_zoom_display(self):
        self.zoom_label.config(text=f"Zoom: {int(self.image_scale * 100)}%")

    def capture_screenshot(self):
        self.root.iconify()
        overlay = ScreenshotOverlay(self.root)
        self.root.wait_window(overlay)
        
        if overlay.selection:
            x1, y1, x2, y2 = map(int, overlay.selection)
            screenshot = pyautogui.screenshot(region=(x1, y1, x2-x1, y2-y1))
            
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            save_path = SCREENSHOT_DIR / f"screenshot_{timestamp}.png"
            screenshot.save(save_path)
            
            self.process_image(save_path)
            self.update_status(f"Screenshot saved to: {save_path}")
        
        self.root.deiconify()

    def upload_image(self):
        file_types = [
            ("Images", "*.png *.jpg *.jpeg *.bmp *.tiff *.webp"),
            ("All Files", "*.*")
        ]
        
        file_path = filedialog.askopenfilename(
            initialdir=str(BASE_DIR),
            filetypes=file_types
        )
        
        if file_path:
            self.process_image(file_path)
            self.update_status(f"Processed: {Path(file_path).name}")

    def process_image(self, image_path):
        try:
            img = cv2.imread(str(image_path))
            gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
            processed = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY | cv2.THRESH_OTSU)[1]

            self.extracted_text = pytesseract.image_to_string(processed)
            self.txt_display.delete(1.0, tk.END)
            self.txt_display.insert(tk.END, self.extracted_text or "No text detected")
            self.show_preview(image_path)
            self.auto_save_text()

        except Exception as e:
            self.show_error(f"Processing Error: {str(e)}")

    def auto_save_text(self):
        if self.extracted_text.strip():
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            save_path = TEXT_DIR / f"extracted_{timestamp}.txt"
            
            try:
                with open(save_path, 'w', encoding='utf-8') as f:
                    f.write(self.extracted_text)
                self.update_status(f"Auto-saved to: {save_path.name}")
            except Exception as e:
                self.show_error(f"Auto-save Failed: {str(e)}")

    def save_text(self, file_path=None):
        if not self.extracted_text.strip():
            self.show_warning("No text to save!")
            return

        if not file_path:
            file_path = filedialog.asksaveasfilename(
                initialdir=str(TEXT_DIR),
                defaultextension=".txt",
                filetypes=[("Text Files", "*.txt"), ("All Files", "*.*")]
            )

        if file_path:
            try:
                with open(file_path, 'w', encoding='utf-8') as f:
                    f.write(self.extracted_text)
                self.update_status(f"Text saved to: {file_path}")
            except Exception as e:
                self.show_error(f"Save Error: {str(e)}")

    def save_word(self):
        if not self.extracted_text.strip():
            self.show_warning("No text to save!")
            return

        file_path = filedialog.asksaveasfilename(
            initialdir=str(TEXT_DIR),
            defaultextension=".docx",
            filetypes=[("Word Documents", "*.docx"), ("All Files", "*.*")]
        )

        if file_path:
            try:
                doc = Document()
                doc.add_paragraph(self.extracted_text)
                doc.save(file_path)
                self.update_status(f"Word document saved to: {file_path}")
            except Exception as e:
                self.show_error(f"Word Save Error: {str(e)}")

    def copy_clipboard(self):
        if self.extracted_text.strip():
            pyperclip.copy(self.extracted_text)
            self.update_status("Text copied to clipboard")
        else:
            self.show_warning("No text to copy!")

    def load_config(self):
        try:
            if CONFIG_FILE.exists():
                with open(CONFIG_FILE, 'r') as f:
                    tesseract_path = f.read().strip()
                    if Path(tesseract_path).exists():
                        pytesseract.pytesseract.tesseract_cmd = tesseract_path
        except Exception as e:
            self.show_error(f"Config Load Error: {str(e)}")

    def configure_tesseract(self):
        file_path = filedialog.askopenfilename(
            title="Select Tesseract Executable",
            initialdir=str(BASE_DIR),
            filetypes=[("Executable Files", "*.exe")]
        )
        
        if file_path and Path(file_path).exists():
            pytesseract.pytesseract.tesseract_cmd = file_path
            try:
                with open(CONFIG_FILE, 'w') as f:
                    f.write(file_path)
                self.update_status(f"Tesseract configured: {file_path}")
            except Exception as e:
                self.show_error(f"Config Save Error: {str(e)}")

    def update_status(self, message):
        self.status.config(text=message)

    def show_error(self, message):
        messagebox.showerror("Error", message)
        self.update_status("Error occurred - see message box")

    def show_warning(self, message):
        messagebox.showwarning("Warning", message)

    def on_close(self):
        keyboard.unhook_all()
        self.root.destroy()

if __name__ == "__main__":
    if sys.platform.startswith('win'):
        default_tesseract = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
        if Path(default_tesseract).exists():
            pytesseract.pytesseract.tesseract_cmd = default_tesseract

    root = tk.Tk()
    app = BifrostApp(root)
    root.mainloop()